from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


output = Path("/home/ubuntu/luma-foundry/Operator2_EmberSignal_Handoff_2026-09-06.docx")
document = Document()
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

title = document.add_heading("Operator A → Operator 2 — Ember Signal Handoff", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = document.add_paragraph("KanPi Foundry | Read-only status handoff | 2026-09-06")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].italic = True

document.add_heading("Purpose", level=1)
document.add_paragraph(
    "This handoff records the current verified Ember Signal evidence endpoint so Operator 2 can assess product, compliance, or integration follow-up without modifying the existing IP Compliance Register. This document does not grant access or alter Drive permissions."
)

document.add_heading("Verified live endpoint", level=1)
table = document.add_table(rows=1, cols=2)
table.style = "Table Grid"
header = table.rows[0].cells
header[0].text = "Item"
header[1].text = "Verified state"
rows = [
    ("Public API", "https://kanpi1inburma.manus.space/api/signals"),
    ("Response state", "HTTP success with JSON status: ready"),
    ("Evidence source", "Credential-safe local evidence document in the KanPi project"),
    ("Evidence record", "My First Evidence | status: observed | source: Local export | observedAt: 2026-09-05"),
    ("top_brands", "[] because the current evidence schema has no verified brand-ranking data; no rankings are inferred"),
]
for label, value in rows:
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value

document.add_heading("Current implementation", level=1)
document.add_paragraph(
    "The application is an Express + tRPC deployment, not Flask. GET /api/signals reads the bounded evidence schema, rejects credential-like content, and returns empty top_brands rather than inventing values. The endpoint is publicly deployed and verified with the supplied record."
)

document.add_heading("Production boundary and requested follow-up", level=1)
document.add_paragraph(
    "The current evidence is a test record that confirms the ingestion path; it is not commercial-clearance evidence. Before a production evidence workflow is declared complete, replace test data with redacted factual records, rotate the previously exposed external API credential, move evidence storage from local filesystem to persistent storage, test the authenticated campaign-brief path, and complete browser-level protected-panel verification."
)

document.add_heading("No changes made to shared compliance records", level=1)
document.add_paragraph(
    "This handoff intentionally does not edit Product Evidence Matrix, Source Commit Register, Asset Licence Register, or Drive sharing settings. Any later compliance status update should be evidence-backed and owner-approved."
)

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)

document.save(output)
print(output)
